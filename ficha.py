"""Gera a FICHA DE ORÇAMENTO igual ao modelo padrão (modelo_ficha.xlsx: abas FICHA e MONT).

FICHA = uma folha por peça (matéria-prima, terceiros, processo, ferramental, ensaios).
MONT  = folha do conjunto: soma as peças x quantidade e inclui a montagem.
Os totais ficam como fórmulas, então dá para ajustar na própria planilha depois.
"""
import io
import re
from copy import deepcopy
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import calculos as calc

VERSAO = "2026-10-31.3"  # os 4 arquivos .py precisam estar na mesma versão

MODELO = Path(__file__).parent / "modelo_ficha.xlsx"

# Grupos de máquina das operações do ABS (iguais aos da ficha preenchida)
GRUPOS_ABS = calc.GRUPOS_ABS
GRUPO_MONTAGEM = calc.GRUPO_MONTAGEM

# Quantas linhas o modelo comporta em cada bloco da FICHA
CAP_MP, CAP_TERC, CAP_OPS, CAP_FERR, CAP_ENS, CAP_OBS = 4, 3, 15, 7, 3, 5
CAP_MONT = 18  # componentes na folha MONT


# -----------------------------------------------------------------------------
# utilidades
# -----------------------------------------------------------------------------

def _txt(v):
    return "" if v is None or (not isinstance(v, str) and pd.isna(v)) else str(v).strip()


def _num(v, padrao=0.0):
    try:
        x = float(v)
        return padrao if x != x else x
    except (TypeError, ValueError):
        return padrao


def _codigo(v):
    """Código como número quando só tem dígitos (como na ficha), senão texto."""
    t = _txt(v)
    if t.endswith(".0"):
        t = t[:-2]
    return int(t) if t.isdigit() else (t or None)


def _virg(v, casas):
    return f"{v:.{casas}f}".replace(".", ",")


def _w(ws, coord, valor):
    """Escreve na célula; se ela estiver dentro de um bloco mesclado, escreve no canto do bloco."""
    for faixa in ws.merged_cells.ranges:
        if coord in faixa:
            coord = faixa.start_cell.coordinate
            break
    ws[coord] = valor


def _nome_aba(codigo, usados):
    base = re.sub(r"[\[\]:*?/\\]", "-", _txt(codigo))[:28] or "PECA"
    nome, n = base, 2
    while nome in usados or nome in ("FICHA", "MONT"):
        nome, n = f"{base}-{n}", n + 1
    return nome


# -----------------------------------------------------------------------------
# componentes (a partir dos orçamentos salvos)
# -----------------------------------------------------------------------------

def _terceiro_pintura(pintura):
    return {"codigo": "", "desc": f'PINTURA {pintura["Tinta"]} ({int(pintura["Lados"])} lado(s))', "qtd": 1,
            "un": "PC", "fornecedor": "", "obs": "", "preco": _num(pintura["Custo Total (R$)"])}


def componente_abs(par, row, pintura=None):
    """Peça ABS (aba ABS salva) -> dados da ficha. Preço de chapa fica em branco (material do cliente)."""
    k = par.k
    cod_chapa = str(_codigo(row.get("Código Chapa")) or "")
    achada = par.chapas[par.chapas["Código"].astype(str) == cod_chapa]
    if cod_chapa and not achada.empty:
        x = achada.iloc[0]
        menor, maior = sorted([_num(x["Comprimento (mm)"]), _num(x["Largura (mm)"])])
        desc = (f'{_txt(x["Descrição"])} {_virg(_num(x["Espessura (mm)"]), 2)} x {menor:.0f} x {maior:.0f}mm')
    else:
        desc = f'CHAPA {_txt(row.get("Polímero"))} {_txt(row.get("Cor"))}'.strip()

    ops = []
    for nome, chave, c_setup, c_tempo in (("MOLDAR PEÇA", "Moldar", "Tempo Prep (min)", "Tempo Termo (min)"),
                                          ("DESTACAR PEÇA", "Destacar", None, "Tempo Destacar (min)"),
                                          ("RECORTAR PEÇA", "Recortar", None, "Tempo Recortar (min)"),
                                          ("DAR ACABAMENTO", "Acabamento", None, "Tempo Acabamento (min)")):
        cc_padrao = k[calc.K_ABS_CC_TERMO] if nome == "MOLDAR PEÇA" else k[calc.K_ABS_CC_ACAB]
        rot = _txt(row.get(f"Grupo {chave}"))
        grupo, desc_grupo = par.separar_grupo(rot) if rot else (GRUPOS_ABS[nome], par.grupo(GRUPOS_ABS[nome]))
        ops.append({"desc": nome, "cc": int(_num(row.get(f"C.C. {chave}"), cc_padrao)), "grupo": grupo,
                    "grupo_desc": desc_grupo, "setup": _num(row.get(c_setup)) if c_setup else 0,
                    "tempo": _num(row.get(c_tempo)), "homens": 1})

    return {
        "codigo": _txt(row["Código Item"]), "descricao": _txt(row.get("Descrição")), "qtd": 1, "tipo": "ABS",
        "lote": _num(row.get("Lote Econômico (pç)"), k[calc.K_LOTE]),
        "mp": [{"codigo": cod_chapa, "desc": desc, "qtd": _num(row.get("Peso Bruto (kg)")), "un": "KG",
                "preco": None, "obs": f'BLANK {_txt(row.get("Blank (mm)")).replace(" x ", " X ")}'},
               {"obs": f'peso liq. {_virg(_num(row.get("Peso Líquido (kg)")), 3)} kg'}],
        "terc": [_terceiro_pintura(pintura)] if pintura is not None else [],
        "ops": ops, "ferram": [], "ensaios": [],
        "obs": [f'{_txt(row.get("Polímero"))} {_txt(row.get("Cor"))} | Conceito: {_txt(row.get("Conceito"))} | '
                f'{_txt(row.get("Tipo Corte"))}'],
        "avisos": [],
    }


def componente_insertos(par, row, detalhe, pintura=None):
    """Peça de insertos (aba Insertos salva + detalhe) -> dados da ficha."""
    cod = _txt(row["Código Item"])
    det = detalhe[detalhe["Código Item"].astype(str) == cod] if not detalhe.empty else detalhe
    comp = {"codigo": cod, "descricao": _txt(row.get("Descrição")), "qtd": 1, "tipo": "Insertos",
            "lote": _num(row.get("Lote Econômico (pç)"), par.k[calc.K_LOTE]),
            "mp": [], "terc": [], "ops": [], "ferram": [], "ensaios": [], "obs": [], "avisos": []}
    for _, x in det.iterrows():
        bloco = x["Bloco"]
        if bloco == "Matéria-prima":
            comp["mp"].append({"codigo": _txt(x.get("Código")), "desc": _txt(x["Descrição"]), "qtd": _num(x["Qtd / Tempo"]),
                               "un": _txt(x.get("Un")), "preco": _num(x["Preço / Taxa (R$)"]), "obs": ""})
        elif bloco == "Terceiros":
            comp["terc"].append({"codigo": "", "desc": _txt(x["Descrição"]), "qtd": _num(x["Qtd / Tempo"]), "un": "",
                                 "fornecedor": _txt(x.get("Fornecedor")), "obs": "", "preco": _num(x["Preço / Taxa (R$)"])})
        elif bloco == "Operação":
            if pd.isna(x.get("C.C.")):
                comp["avisos"].append(f"{cod}: operação '{x['Descrição']}' salva numa versão antiga (sem centro de custo). "
                                      "Salve o orçamento de novo.")
                continue
            comp["ops"].append({"desc": _txt(x["Descrição"]), "cc": int(float(x["C.C."])), "grupo": _txt(x.get("Grupo Máquina")),
                                "grupo_desc": _txt(x.get("Descrição Grupo")) or par.grupo(_txt(x.get("Grupo Máquina"))),
                                "setup": _num(x.get("Setup (min)")), "tempo": _num(x["Qtd / Tempo"]),
                                "homens": _num(x.get("Nº Homens"), 1)})
        elif bloco == "Ferramental":
            comp["ferram"].append({"qtd": 1, "desc": _txt(x["Descrição"]), "valor": _num(x["Total (R$)"])})
        elif bloco == "Ensaio":
            comp["ensaios"].append({"desc": _txt(x["Descrição"]), "valor": _num(x["Total (R$)"])})
    if pintura is not None:
        comp["terc"].append(_terceiro_pintura(pintura))
    return comp


# -----------------------------------------------------------------------------
# custos (mesma lógica das fórmulas da ficha)
# -----------------------------------------------------------------------------

def custos_peca(par, comp):
    lote = max(_num(comp["lote"], 1), 1)
    mp = sum(_num(m.get("qtd")) * _num(m.get("preco")) for m in comp["mp"])
    ext = sum(_num(t.get("qtd")) * _num(t.get("preco")) for t in comp["terc"])
    mo = sum(par.taxa(o["cc"]) * _num(o["tempo"]) / 60 * _num(o["homens"], 1) for o in comp["ops"])
    setup = sum(par.taxa(o["cc"]) * _num(o["setup"]) / lote / 60 for o in comp["ops"])
    return {"Matéria-prima": mp, "Externo": ext, "Mão de obra": mo, "Set-up": setup, "Total": mp + ext + mo + setup}


def custo_montagem(par, mont, lote):
    if not mont:
        return 0.0, 0.0
    taxa = par.taxa(mont["cc"])
    return taxa * _num(mont["setup"]) / max(_num(lote, 1), 1) / 60, taxa * _num(mont["tempo"]) / 60 * _num(mont["homens"], 1)


def totais_conjunto(par, comps, mont, lote):
    """Mesmas somas da folha MONT (set-up não multiplica pela quantidade, como no modelo)."""
    t = {"Matéria-prima": 0.0, "Externo": 0.0, "Mão de obra": 0.0, "Set-up": 0.0}
    for c in comps:
        x = custos_peca(par, c)
        t["Matéria-prima"] += x["Matéria-prima"] * c["qtd"]
        t["Externo"] += x["Externo"] * c["qtd"]
        t["Mão de obra"] += x["Mão de obra"] * c["qtd"]
        t["Set-up"] += x["Set-up"]
    m_setup, m_mo = custo_montagem(par, mont, lote)
    t["Set-up"] += m_setup
    t["Mão de obra"] += m_mo
    t["Total"] = sum(t.values())
    return t


# -----------------------------------------------------------------------------
# preenchimento do modelo
# -----------------------------------------------------------------------------

def _preencher_peca(ws, comp, par):
    avisos = list(comp.get("avisos", []))
    _w(ws, "C2", comp["descricao"] or None)
    _w(ws, "H2", _codigo(comp["codigo"]))
    _w(ws, "K2", comp["lote"])

    if len(comp["mp"]) > CAP_MP:
        avisos.append(f'{comp["codigo"]}: só cabem {CAP_MP} linhas de matéria-prima na ficha; as demais foram omitidas.')
    for i in range(CAP_MP):
        r = 6 + i
        m = comp["mp"][i] if i < len(comp["mp"]) else {}
        _w(ws, f"A{r}", _codigo(m.get("codigo")))
        _w(ws, f"B{r}", m.get("desc") or None)
        _w(ws, f"E{r}", m.get("qtd"))
        _w(ws, f"F{r}", m.get("un") or None)
        _w(ws, f"G{r}", m.get("preco"))
        _w(ws, f"I{r}", m.get("obs") or None)

    if len(comp["terc"]) > CAP_TERC:
        avisos.append(f'{comp["codigo"]}: só cabem {CAP_TERC} operações externas na ficha; as demais foram omitidas.')
    for i in range(CAP_TERC):
        r = 15 + i
        t = comp["terc"][i] if i < len(comp["terc"]) else {}
        _w(ws, f"A{r}", _codigo(t.get("codigo")))
        _w(ws, f"B{r}", t.get("desc") or None)
        _w(ws, f"C{r}", t.get("qtd"))
        _w(ws, f"D{r}", t.get("un") or None)
        _w(ws, f"E{r}", t.get("fornecedor") or None)
        _w(ws, f"H{r}", t.get("obs") or None)
        _w(ws, f"J{r}", t.get("preco"))

    if len(comp["ops"]) > CAP_OPS:
        avisos.append(f'{comp["codigo"]}: só cabem {CAP_OPS} operações na ficha; as demais foram omitidas.')
    for i in range(CAP_OPS):
        r = 23 + i
        o = comp["ops"][i] if i < len(comp["ops"]) else None
        if o is None:
            continue
        taxa = par.taxa(o["cc"])
        _w(ws, f"B{r}", o["desc"])
        _w(ws, f"C{r}", o["cc"])
        _w(ws, f"D{r}", o["grupo"] or None)
        _w(ws, f"E{r}", o.get("grupo_desc") or par.grupo(o["grupo"]) or None)
        _w(ws, f"G{r}", o["setup"])
        _w(ws, f"H{r}", o["tempo"])
        _w(ws, f"I{r}", o["homens"])
        _w(ws, f"J{r}", f"={taxa!r}*G{r}/$K$2/60")
        _w(ws, f"K{r}", f"={taxa!r}*H{r}/60*I{r}")

    for i, f in enumerate(comp["ferram"][:CAP_FERR]):
        _w(ws, f"A{45 + i}", f.get("qtd", 1))
        _w(ws, f"B{45 + i}", f["desc"])
        _w(ws, f"E{45 + i}", f["valor"])
    for i, e in enumerate(comp["ensaios"][:CAP_ENS]):
        _w(ws, f"A{64 + i}", e["desc"])
        _w(ws, f"K{64 + i}", e["valor"])
    for i, texto in enumerate(comp["obs"][:CAP_OBS]):
        _w(ws, f"A{70 + i}", texto)
    return avisos


def _preencher_mont(ws, cab, comps, nomes, mont, par):
    _w(ws, "I1", cab.get("numero") or None)
    _w(ws, "B2", cab.get("cliente") or None)
    _w(ws, "H2", _codigo(cab.get("codigo")))
    _w(ws, "C3", cab.get("descricao") or None)
    _w(ws, "H3", cab.get("consumo") or None)
    _w(ws, "K3", cab.get("lote"))
    _w(ws, "B4", cab.get("similar") or None)

    for i, (comp, nome) in enumerate(zip(comps[:CAP_MONT], nomes)):
        r, ref = 9 + i, "'" + nome.replace("'", "''") + "'"
        _w(ws, f"A{r}", f"={ref}!$H$2")
        _w(ws, f"B{r}", f"={ref}!$C$2")
        _w(ws, f"C{r}", comp["qtd"])
        _w(ws, f"D{r}", "UN")
        _w(ws, f"E{r}", f"={ref}!$K$10")
        _w(ws, f"G{r}", f"={ref}!$K$18")
        _w(ws, f"I{r}", f"={ref}!$J$38")
        _w(ws, f"J{r}", f"={ref}!$K$38")

    if mont:
        taxa = par.taxa(mont["cc"])
        _w(ws, "B57", "MONTAR CONJUNTO")
        _w(ws, "C57", mont["cc"])
        _w(ws, "D57", mont["grupo"])
        _w(ws, "E57", mont.get("grupo_desc") or par.grupo(mont["grupo"]) or None)
        _w(ws, "G57", mont["setup"])
        _w(ws, "H57", mont["tempo"])
        _w(ws, "I57", mont["homens"])
        _w(ws, "J57", f"={taxa!r}*G57/$K$3/60")
        _w(ws, "K57", f"={taxa!r}*(H57/60*I57)")


def gerar(par, cab, comps, mont=None, incluir_mont=True):
    """Devolve (bytes do .xlsx, avisos). comps = lista de componentes (ver componente_abs/insertos)."""
    wb = load_workbook(MODELO)
    modelo = wb["FICHA"]
    avisos, nomes = [], []

    for comp in comps:
        comp = dict(comp, lote=cab.get("lote") or comp["lote"])
        nome = _nome_aba(comp["codigo"], nomes)
        nomes.append(nome)
        ws = wb.copy_worksheet(modelo)
        ws.title = nome
        for dv in modelo.data_validations.dataValidation:       # lista do "Nº de homens" etc.
            ws.add_data_validation(deepcopy(dv))
        for faixa in modelo.conditional_formatting:
            for regra in faixa.rules:
                ws.conditional_formatting.add(str(faixa.sqref), deepcopy(regra))
        avisos += _preencher_peca(ws, comp, par)
    wb.remove(modelo)

    if incluir_mont and comps:
        if len(comps) > CAP_MONT:
            avisos.append(f"A folha MONT comporta {CAP_MONT} componentes; os demais ficaram só nas folhas das peças.")
        _preencher_mont(wb["MONT"], cab, comps, nomes, mont, par)
        wb.move_sheet("MONT", offset=len(wb.sheetnames) - 1)
    else:
        wb.remove(wb["MONT"])

    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is wb.worksheets[0]
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue(), avisos
